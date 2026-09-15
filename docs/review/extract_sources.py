"""Read supporting artifacts without changing originals; retain locators and hashes."""
from pathlib import Path
import hashlib
import json
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs/review/source-extracts"
OUT.mkdir(parents=True, exist_ok=True)
manifest = []
for path in sorted(ROOT.rglob("*")):
    if "docs" in path.relative_to(ROOT).parts or path.suffix.lower() not in {".pdf", ".docx", ".xlsx"}:
        continue
    rel = path.relative_to(ROOT).as_posix()
    lines = [f"SOURCE: {rel}"]
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(path)
        lines.append(f"PAGES: {len(reader.pages)}")
        for number, page in enumerate(reader.pages, 1):
            lines.extend([f"\n--- PDF PAGE {number} ---", page.extract_text() or "[NO EXTRACTABLE TEXT]"])
    elif path.suffix.lower() == ".docx":
        doc = Document(path)
        lines.extend(f"P{i}: {p.text}" for i, p in enumerate(doc.paragraphs, 1) if p.text.strip())
        for i, table in enumerate(doc.tables, 1):
            lines.append(f"TABLE {i}")
            lines.extend(f"ROW {j}: " + " | ".join(c.text for c in row.cells) for j, row in enumerate(table.rows, 1))
    else:
        book = load_workbook(path, read_only=True, data_only=False)
        for sheet in book:
            lines.append(f"SHEET: {sheet.title}")
            for row in sheet:
                values = [f"{c.coordinate}={c.value}" for c in row if c.value is not None]
                if values:
                    lines.append(" | ".join(values))
        book.close()
    name = rel.replace("/", "__") + ".txt"
    (OUT / name).write_text("\n".join(lines), encoding="utf-8")
    manifest.append({"path": rel, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "extraction": f"docs/review/source-extracts/{name}"})
(OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"sources": len(manifest), "manifest": str(OUT / "manifest.json")}, ensure_ascii=False))
